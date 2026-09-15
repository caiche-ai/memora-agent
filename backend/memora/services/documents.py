from __future__ import annotations

import io
import re
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError
from pypdf import PdfReader

TEXT_EXTENSIONS = {".txt", ".md", ".markdown"}
PRESENTATION_EXTENSIONS = {".pptx"}
SUPPORTED_EXTENSIONS = {".pdf", *TEXT_EXTENSIONS, *PRESENTATION_EXTENSIONS}


def extract_pdf(data: bytes) -> dict[str, Any]:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as error:
            raise ValueError("PDF 已加密，无法解析") from error
    pages = []
    for page_number, page in enumerate(reader.pages, 1):
        text = re.sub(r"\s+", " ", page.extract_text() or "").strip()
        pages.append({"page_number": page_number, "text": text})
    return {"pages": pages, "text": "\n\n".join(item["text"] for item in pages)}


def _normalize_presentation_text(value: str) -> str:
    lines = [re.sub(r"\s+", " ", line).strip() for line in value.replace("\v", "\n").splitlines()]
    return "\n".join(line for line in lines if line)


def _shape_text_blocks(shapes: Any) -> list[str]:
    blocks: list[str] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            blocks.extend(_shape_text_blocks(shape.shapes))
            continue
        if getattr(shape, "has_table", False):
            rows = []
            for row in shape.table.rows:
                cells = [_normalize_presentation_text(cell.text) for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append("\n".join(rows))
            continue
        if getattr(shape, "has_text_frame", False):
            text = _normalize_presentation_text(shape.text)
            if text:
                blocks.append(text)
    return blocks


def extract_pptx(data: bytes) -> dict[str, Any]:
    try:
        presentation = Presentation(io.BytesIO(data))
    except (PackageNotFoundError, ValueError, KeyError) as error:
        raise ValueError("PPTX 文件损坏或格式无效") from error

    pages = []
    for page_number, slide in enumerate(presentation.slides, 1):
        blocks = _shape_text_blocks(slide.shapes)
        if getattr(slide, "has_notes_slide", False):
            notes_frame = slide.notes_slide.notes_text_frame
            notes = _normalize_presentation_text(notes_frame.text if notes_frame else "")
            if notes:
                blocks.append(f"演讲者备注：\n{notes}")
        pages.append({"page_number": page_number, "text": "\n\n".join(blocks)})
    return {"pages": pages, "text": "\n\n".join(item["text"] for item in pages)}


def chunk_pages(
    pages: list[dict[str, Any]], max_length: int = 900, overlap: int = 120
) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    index = 0
    for page in pages:
        if not page["text"]:
            continue
        heading = ""
        sections: list[tuple[str, str]] = []
        section_lines: list[str] = []
        for line in str(page["text"]).splitlines():
            heading_match = re.match(r"^\s{0,3}#{1,6}\s+(.+?)\s*$", line)
            if heading_match:
                if section_lines:
                    sections.append((heading, "\n".join(section_lines)))
                    section_lines = []
                heading = heading_match.group(1).strip()
            else:
                section_lines.append(line)
        if section_lines or not sections:
            sections.append((heading, "\n".join(section_lines)))
        current = ""
        current_heading = ""
        for section_heading, section_text in sections:
            paragraphs = [item for item in re.split(r"(?<=[。！？.!?])\s*", section_text) if item]
            for paragraph in paragraphs:
                if len(current) + len(paragraph) > max_length and current:
                    chunks.append(
                        {
                            "index": index,
                            "page_number": page["page_number"],
                            "heading": current_heading,
                            "content": current.strip(),
                        }
                    )
                    index += 1
                    current = current[-overlap:] + paragraph
                else:
                    current += paragraph
                current_heading = section_heading or current_heading
            if current.strip():
                chunks.append(
                    {
                        "index": index,
                        "page_number": page["page_number"],
                        "heading": current_heading,
                        "content": current.strip(),
                    }
                )
                index += 1
                current = ""
                current_heading = ""
    return chunks


def decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("gb18030", errors="replace")


def extract_knowledge_file(filename: str, data: bytes) -> dict[str, Any]:
    extension = f".{filename.lower().rsplit('.', 1)[-1]}" if "." in filename else ""
    if extension == ".pdf":
        result = extract_pdf(data)
        result["format"] = "pdf"
        return result
    if extension in PRESENTATION_EXTENSIONS:
        result = extract_pptx(data)
        result["format"] = "pptx"
        return result
    if extension in TEXT_EXTENSIONS:
        text = decode_text(data).strip()
        return {
            "pages": [{"page_number": 1, "text": text}],
            "text": text,
            "format": extension.removeprefix("."),
        }
    raise ValueError("不支持的文件类型；目前支持 PDF、PPTX、TXT 和 Markdown")
