from __future__ import annotations

import asyncio
import io
import sqlite3
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from memora.config import load_config
from memora.services.documents import chunk_pages, decode_text, extract_knowledge_file
from memora.services.meeting import analyze_meeting, heuristic_analysis, meeting_memories
from memora.services.ppt import create_presentation, is_ppt_request
from memora.store import Store, tokenize


def test_rerank_config_reuses_llm_api_key(monkeypatch) -> None:
    monkeypatch.setenv("LLM_API_KEY", "shared-key")
    monkeypatch.setenv("RERANK_API_KEY", "")
    monkeypatch.setenv("RERANK_MODEL", "gte-rerank-v2")
    monkeypatch.setenv("RERANK_TOP_N", "12")

    settings = load_config().rerank

    assert settings.api_key == "shared-key"
    assert settings.model == "gte-rerank-v2"
    assert settings.top_n == 12


def test_store_conversations_documents_and_memories() -> None:
    store = Store(":memory:")
    conversation = store.create_conversation()
    store.add_message(conversation["id"], "user", "项目交付计划是什么？")
    document = store.add_document(
        conversation_id=conversation["id"],
        name="plan.pdf",
        mime_type="application/pdf",
        kind="knowledge",
        text="项目将在九月交付",
    )
    store.add_chunks(
        document["id"], [{"index": 0, "page_number": 2, "content": "星河项目计划在九月完成交付。"}]
    )
    store.add_memory(
        {
            "type": "person",
            "subject": "张三",
            "content": "负责星河项目",
            "source_type": "chat",
            "source_id": 1,
        }
    )
    assert len(store.list_messages(conversation["id"])) == 1
    assert store.search_chunks(conversation["id"], tokenize("星河项目什么时候交付"))[0]["page_number"] == 2
    assert store.search_memories("张三负责什么")[0]["subject"] == "张三"
    store.close()


def test_existing_meetings_migrate_to_default_project(tmp_path: Path) -> None:
    database = tmp_path / "legacy.db"
    connection = sqlite3.connect(database)
    connection.executescript(
        """
        CREATE TABLE documents (
          id INTEGER PRIMARY KEY AUTOINCREMENT, conversation_id INTEGER,
          name TEXT NOT NULL, mime_type TEXT NOT NULL, kind TEXT NOT NULL,
          status TEXT NOT NULL DEFAULT 'ready', text_content TEXT NOT NULL DEFAULT '',
          metadata_json TEXT NOT NULL DEFAULT '{}', created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE meetings (
          id INTEGER PRIMARY KEY AUTOINCREMENT, document_id INTEGER NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
          title TEXT NOT NULL, summary TEXT NOT NULL, participants_json TEXT NOT NULL DEFAULT '[]',
          dates_json TEXT NOT NULL DEFAULT '[]', locations_json TEXT NOT NULL DEFAULT '[]',
          topics_json TEXT NOT NULL DEFAULT '[]', risks_json TEXT NOT NULL DEFAULT '[]',
          decisions_json TEXT NOT NULL DEFAULT '[]', created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        INSERT INTO documents(name,mime_type,kind,text_content) VALUES ('legacy.txt','text/plain','meeting','旧会议');
        INSERT INTO meetings(document_id,title,summary) VALUES (1,'旧会议','迁移测试');
        """
    )
    connection.commit()
    connection.close()

    store = Store(database)
    projects = store.list_projects()
    assert projects[0]["name"] == "默认项目"
    assert "project_id" in {row["name"] for row in store._all("PRAGMA table_info(conversations)")}
    assert "project_id" in {row["name"] for row in store._all("PRAGMA table_info(documents)")}
    assert store.list_meetings(projects[0]["id"])[0]["title"] == "旧会议"
    store.close()


def test_meeting_analysis_extracts_risks_todos_and_memories() -> None:
    text = """星河项目周例会
时间：2026年8月25日
地点：三号会议室
张三：测试环境还没有准备完成，存在延期风险。
李四：由李四在8月27日前完成环境部署。
决定保持验收节点不变。"""
    result = heuristic_analysis(text, "meeting.txt")
    assert "张三" in result["participants"]
    assert "2026年8月25日" in result["dates"]
    assert result["risks"]
    assert result["todos"]
    assert result["decisions"]
    assert any(item["type"] == "risk" for item in meeting_memories(result, 1))


def test_meeting_analysis_receives_project_background(monkeypatch) -> None:
    captured: list[dict[str, str]] = []

    async def fake_complete_json(messages, fallback):
        captured.extend(messages)
        return fallback

    monkeypatch.setattr("memora.services.meeting.complete_json", fake_complete_json)
    result = asyncio.run(
        analyze_meeting(
            "本周确认验收安排。",
            "meeting.txt",
            [{"name": "产品规范.pdf", "content": "项目验收必须完成安全评审。"}],
        )
    )
    assert result["title"]
    assert "产品规范.pdf" in captured[1]["content"]
    assert "项目验收必须完成安全评审" in captured[1]["content"]


def test_document_chunking_and_text_decoding() -> None:
    chunks = chunk_pages([{"page_number": 3, "text": "这是第一段。" * 100 + "这是结尾。"}], 120, 20)
    assert len(chunks) > 2
    assert all(item["page_number"] == 3 for item in chunks)
    assert decode_text("普通文本".encode()) == "普通文本"


def test_pptx_extraction_preserves_slides_tables_and_notes() -> None:
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "项目总览"
    first.placeholders[1].text = "一期计划九月交付"
    first.notes_slide.notes_text_frame.text = "交付前必须完成验收"

    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "风险清单"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    table.cell(0, 0).text = "风险"
    table.cell(0, 1).text = "负责人"
    table.cell(1, 0).text = "接口延期"
    table.cell(1, 1).text = "张三"

    buffer = io.BytesIO()
    presentation.save(buffer)
    extracted = extract_knowledge_file("项目汇报.pptx", buffer.getvalue())

    assert extracted["format"] == "pptx"
    assert len(extracted["pages"]) == 2
    assert extracted["pages"][0]["page_number"] == 1
    assert "项目总览" in extracted["pages"][0]["text"]
    assert "一期计划九月交付" in extracted["pages"][0]["text"]
    assert "演讲者备注" in extracted["pages"][0]["text"]
    assert "交付前必须完成验收" in extracted["pages"][0]["text"]
    assert "风险 | 负责人" in extracted["pages"][1]["text"]
    assert "接口延期 | 张三" in extracted["pages"][1]["text"]
    chunks = chunk_pages(extracted["pages"])
    assert {item["page_number"] for item in chunks} == {1, 2}


def test_ppt_request_and_generation(tmp_path: Path) -> None:
    assert is_ppt_request("帮我制作一份项目汇报 PPT")
    result = asyncio.run(create_presentation("制作一份项目汇报 PPT", tmp_path))
    content = Path(result["file_path"]).read_bytes()
    assert content[:2] == b"PK"
    assert len(content) > 10_000
